"""Проверки телефонов, дат, дублей и полного сохранения строк в отчёте."""
from collections import Counter
from datetime import date, datetime
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_WINDOWS_1900, to_excel

import clean_zayavki as app


class CleaningTests(unittest.TestCase):
    def test_phone_formats_and_notes(self):
        for raw in [89172405219, 89172405219.0, '+7 (917) 240-52-19',
                    '9172405219', 'тел.: 8 917 240 52 19 (после 18:00)']:
            with self.subTest(raw=raw):
                self.assertEqual(app.normalize_phone(raw), ('+79172405219', ''))
        for raw in ['', None, 'abc', '123', '0000000000', '+12025550123',
                    '+89172405219', '123456789012', 89172405219.5]:
            with self.subTest(raw=raw):
                self.assertFalse(app.normalize_phone(raw)[0])
                self.assertTrue(app.normalize_phone(raw)[1])
        self.assertFalse(app.normalize_phone('+79172405219; +79162451783')[0])

    def test_dates_are_checked_not_guessed(self):
        epoch = CALENDAR_WINDOWS_1900
        for raw in ['09/03/2026', '03.09.2026', '3 сентября 2026',
                    '2026/09/03', '03.09.26', date(2026, 9, 3),
                    datetime(2026, 9, 3, 14, 30), to_excel(datetime(2026, 9, 3))]:
            with self.subTest(raw=raw):
                self.assertEqual(app.resolve_date(raw, epoch, (2026, 9), 'auto'),
                                 (date(2026, 9, 3), ''))
        self.assertEqual(app.resolve_date('12 сентября', epoch, (2027, 9), 'auto'),
                         (date(2027, 9, 12), ''))
        self.assertTrue(app.resolve_date('09/03/2026', epoch, None, 'auto')[1])
        self.assertEqual(app.resolve_date('09/03/2026', epoch, None, 'dmy')[0], date(2026, 3, 9))
        self.assertEqual(app.resolve_date('09/03/2026', epoch, None, 'mdy')[0], date(2026, 9, 3))
        self.assertIsNone(app.resolve_date('09/16/2026', epoch, None, 'dmy')[0])
        self.assertTrue(app.resolve_date('19.09.2062', epoch, (2026, 9), 'auto')[1])
        self.assertIsNone(app.resolve_date('31.09.2026', epoch, None, 'auto')[0])
        self.assertIsNone(app.resolve_date('29.02.2026', epoch, None, 'auto')[0])
        self.assertEqual(app.resolve_date('29.02.2028', epoch, None, 'auto')[0], date(2028, 2, 29))

    def test_names_and_identity_conflicts(self):
        self.assertEqual(app.normalize_name('  СМИРНОВ\u00a0   олег\u200b '), 'Смирнов Олег')
        self.assertEqual(app.normalize_name('Пeтров Алексей'), 'Петров Алексей')
        self.assertEqual(app.normalize_name('John SMITH'), 'John Smith')
        self.assertTrue(app.names_compatible('К. Белоусов', 'Белоусов Кирилл'))
        self.assertTrue(app.names_compatible('Дарья Королева', 'Королёва Дарья'))
        self.assertFalse(app.names_compatible('Белов Никита', 'Морозова Елена'))

    def test_other_workbook_preserves_every_row(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            source = folder / 'другие заявки.xlsx'
            book = Workbook()
            sheet = book.active
            sheet.title = 'Данные'
            sheet.append(['Источник', 'Телефон', 'Имя', 'Дата заявки'])
            sheet.append(['=1+1', '89162451783', 'ИВАН ПЕТРОВ', '18.02.2028'])
            sheet.append(['сайт', '+7 916 245 17 83', 'Петров Иван', '02/18/2028'])
            sheet.append(['email', '89162451783', 'Другой Человек', '19.02.2028'])
            sheet.append(['сайт', 'не знаю', 'АННА', 'хх.02.2028'])
            sheet.append(['email', '89032218840', None, '20.02.2028'])
            sheet.append(['сайт', '89256003112', 'К. Сидоров', '21.02.2028'])
            sheet.append(['email', '+79256003112', 'Сидоров Кирилл', '21.02.2028'])
            sheet.append([None] * 4)
            book.save(source)
            self.assertEqual(app.main([str(source), '--out-dir', str(folder / 'out')]), 0)
            clean = load_workbook(folder / 'out' / 'другие заявки_clean.xlsx')
            report = load_workbook(folder / 'out' / 'другие заявки_problems.xlsx')
            self.assertEqual(clean.active.max_row - 1, 2)
            self.assertEqual(clean.active['A3'].value, 'Сидоров Кирилл')
            self.assertEqual(clean.active['C2'].value, '2028-02-18')
            self.assertEqual(clean.active['D2'].value, '=1+1')
            self.assertEqual(clean.active['D2'].data_type, 's')
            self.assertEqual(report['Аудит'].max_row - 1, 7)
            statuses = Counter(row[2] for row in report['Аудит'].iter_rows(min_row=2, values_only=True))
            self.assertEqual(statuses, {'Чистая': 2, 'Дубликат': 3, 'Проблема': 2})
            clean.close()
            report.close()
            self.assertEqual(app.main([str(source), '--out-dir', str(folder / 'out')]), 1)

    def test_no_consensus_period_does_not_invent_dates(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'many_months.xlsx'
            book = Workbook()
            sheet = book.active
            sheet.append(app.HEADERS)
            for phone, raw_date in [('89162451783', '18.02.2028'), ('89032218840', '18.03.2028'),
                                    ('89256003112', '02/03/2028')]:
                sheet.append(['Иван', phone, raw_date, 'сайт'])
            book.save(source)
            records, period, _, _ = app.load_records(source)
            self.assertIsNone(period)
            self.assertIsNone(records[2].day)
            self.assertIn('Неоднозначная дата', records[2].issues[0])
            self.assertEqual(len(app.select_records(records)), 2)

    def test_worksheet_selection_and_empty_file(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'sheets.xlsx'
            book = Workbook()
            book.active.title = 'Лист 1'
            book.active.append(app.HEADERS)
            book.create_sheet('Лист 2').append(app.HEADERS)
            book.save(source)
            with self.assertRaises(ValueError):
                app.load_records(source)
            records, period, _, _ = app.load_records(source, sheet_name='Лист 2')
            self.assertEqual(records, [])
            self.assertIsNone(period)

    def test_explicit_date_order_is_used_for_period_inference(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'dmy.xlsx'
            book = Workbook()
            sheet = book.active
            sheet.append(app.HEADERS)
            for raw_date in ['01/02/2028', '02/02/2028', '03/02/2028']:
                sheet.append(['Иван', '89162451783', raw_date, 'сайт'])
            book.save(source)
            records, period, _, _ = app.load_records(source, slash_order='dmy')
            self.assertEqual(period, (2028, 2))
            self.assertEqual([r.day.day for r in records], [1, 2, 3])
            self.assertFalse(any(r.issues for r in records))


if __name__ == '__main__':
    unittest.main(verbosity=2)
