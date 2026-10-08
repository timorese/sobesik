#!/usr/bin/env python3
"""Очистка XLSX-заявок. Python 3.10+, openpyxl. Без привязки к Windows/Excel."""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
import math
from pathlib import Path
import re
import sys
import tempfile
import unicodedata
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.utils.datetime import from_excel

HEADERS = ("Имя", "Телефон", "Дата заявки", "Источник")
EMPTY = {"", "-", "—", "нет", "none", "null", "n/a", "не указано"}
MONTHS = {
    "январь": 1, "января": 1, "февраль": 2, "февраля": 2,
    "март": 3, "марта": 3, "апрель": 4, "апреля": 4,
    "май": 5, "мая": 5, "июнь": 6, "июня": 6,
    "июль": 7, "июля": 7, "август": 8, "августа": 8,
    "сентябрь": 9, "сентября": 9, "октябрь": 10, "октября": 10,
    "ноябрь": 11, "ноября": 11, "декабрь": 12, "декабря": 12,
}
LOOKALIKES = str.maketrans(dict(zip("ABCEHKMOPTXYaceopxy", "АВСЕНКМОРТХУасеорху")))
PHONE_RUN = re.compile(r"(?<![0-9])\+?[0-9](?:[0-9\s().\-–—]*[0-9])?(?![0-9])")


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return str(int(value))
    return str(value)


def tidy(value) -> str:
    s = unicodedata.normalize("NFC", text(value))
    s = "".join(c for c in s if unicodedata.category(c) != "Cf")
    return " ".join(s.split())


def normalize_name(value) -> str:
    s = tidy(value)
    if s.casefold() in EMPTY or not any(c.isalpha() for c in s):
        return ""
    def fix(match):
        word = match.group()
        return word.translate(LOOKALIKES) if re.search(r"[А-Яа-яЁё]", word) else word
    s = re.sub(r"[^\W\d_]+", fix, s, flags=re.UNICODE)
    return s.title()


def normalize_phone(value) -> tuple[str, str]:
    s = tidy(value)
    if s.casefold() in EMPTY:
        return "", "Нет телефона"
    if isinstance(value, bool) or (isinstance(value, float) and not value.is_integer()):
        return "", "Нет телефона: некорректное число"
    candidates = []
    for match in PHONE_RUN.finditer(s):
        digits = re.sub(r"[^0-9]", "", match.group())
        if len(digits) == 11 and digits[0] in "78":
            if match.group().startswith("+") and digits[0] != "7":
                continue
            national = digits[1:]
        elif len(digits) == 10 and not match.group().startswith("+"):
            national = digits
        else:
            continue
        if national[0] == "0" or len(set(national)) == 1:
            continue
        candidates.append("+7" + national)
    candidates = list(dict.fromkeys(candidates))
    if len(candidates) == 1:
        return candidates[0], ""
    if len(candidates) > 1:
        return "", "Нет единственного телефона: найдено несколько номеров"
    return "", "Нет телефона: номер не распознан"


def make_date(year: int, month: int, day: int):
    try:
        return date(year + 2000 if 0 <= year < 100 else year, month, day)
    except (ValueError, OverflowError):
        return None


def date_candidates(value, epoch, year: int | None = None) -> list[date]:
    if isinstance(value, datetime):
        return [value.date()]
    if isinstance(value, date):
        return [value]
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            if not math.isfinite(value) or not 1 <= value <= 2958465:
                return []
            converted = from_excel(value, epoch)
            return [converted.date()] if isinstance(converted, datetime) else []
        except (ValueError, OverflowError, TypeError):
            return []
    s = tidy(value).casefold()
    if s in EMPTY:
        return []
    # ISO дата, в том числе ISO дата-время.
    if re.fullmatch(r"[0-9]{4}-[0-9]{2}-[0-9]{2}(?:[t ][0-9]{2}:[0-9]{2}(?::[0-9]{2}(?:\.[0-9]+)?)?(?:z|[+-][0-9]{2}:[0-9]{2})?)?", s):
        try:
            return [datetime.fromisoformat(s.upper().replace("Z", "+00:00")).date()]
        except ValueError:
            return []
    m = re.fullmatch(r"([0-9]{4})[./-]([0-9]{1,2})[./-]([0-9]{1,2})", s)
    if m:
        result = make_date(*map(int, m.groups()))
        return [result] if result else []
    m = re.fullmatch(r"([0-9]{1,2})([./-])([0-9]{1,2})\2([0-9]{2}|[0-9]{4})", s)
    if m:
        a, sep, b, y = m.groups()
        result = [make_date(int(y), int(b), int(a))]
        if sep == "/":
            result.append(make_date(int(y), int(a), int(b)))
        return sorted({d for d in result if d is not None})
    m = re.fullmatch(r"([0-9]{1,2})\s+([а-яё]+)(?:\s+([0-9]{2}|[0-9]{4}))?(?:\s*г\.?)?", s)
    if m and m[2] in MONTHS:
        y = int(m[3]) if m[3] else year
        if y is not None:
            result = make_date(y, MONTHS[m[2]], int(m[1]))
            return [result] if result else []
    return []


def resolve_date(value, epoch, period, slash_order):
    candidates = date_candidates(value, epoch, period[0] if period else None)
    slash = re.fullmatch(r"([0-9]{1,2})/([0-9]{1,2})/([0-9]{2}|[0-9]{4})", tidy(value))
    if slash and slash_order != "auto":
        a, b, y = map(int, slash.groups())
        chosen = make_date(y, b, a) if slash_order == "dmy" else make_date(y, a, b)
        candidates = [chosen] if chosen else []
    if not candidates:
        return None, "Нет даты" if tidy(value).casefold() in EMPTY else "Дата не распознана"
    if len(candidates) > 1:
        in_period = [d for d in candidates if period and (d.year, d.month) == period]
        chosen = in_period[0] if len(in_period) == 1 else None
        if chosen is None:
            return None, "Неоднозначная дата: " + " или ".join(d.isoformat() for d in candidates)
    else:
        chosen = candidates[0]
    if period and (chosen.year, chosen.month) != period:
        return chosen, f"Дата вне периода {period[0]:04d}-{period[1]:02d}"
    return chosen, ""


@dataclass
class Record:
    row: int
    raw: tuple
    name: str
    phone: str
    day: date | None
    source: str
    issues: list[str] = field(default_factory=list)
    kept_row: int | None = None
    status: str = ""


def names_compatible(a: str, b: str) -> bool:
    def tokens(s):
        return re.findall(r"[^\W\d_]+", s.casefold().replace("ё", "е"))
    left, right = tokens(a), tokens(b)
    if len(left) != len(right):
        return False
    # Инициалы сравниваются с первым символом полного имени, порядок не важен.
    def match(rest, available):
        if not rest:
            return True
        first = rest[0]
        for i, other in enumerate(available):
            if first == other or (len(first) == 1 and other.startswith(first)) or (len(other) == 1 and first.startswith(other)):
                if match(rest[1:], available[:i] + available[i + 1:]):
                    return True
        return False
    return match(left, right)


def select_records(records):
    groups = defaultdict(list)
    for r in records:
        if r.phone:
            groups[r.phone].append(r)
        else:
            r.status = "Проблема"
    clean = []
    for group in groups.values():
        def rank(r):
            valid = int(bool(r.name)) + int(r.day is not None and not any("дат" in i.casefold() for i in r.issues))
            full_words = sum(len(t) > 1 for t in re.findall(r"[^\W\d_]+", r.name))
            return (-valid, -full_words, r.day or date.max, r.row)
        kept = min(group, key=rank)
        conflict = any(a.name and b.name and not names_compatible(a.name, b.name)
                       for i, a in enumerate(group) for b in group[i + 1:])
        for r in group:
            r.kept_row = kept.row
            if conflict:
                r.issues.append("Разные имена у одного телефона: требуется проверка")
            if r is not kept:
                r.status = "Дубликат"
                r.issues.append(f"Дубликат телефона; выбрана строка {kept.row}")
            else:
                # Конфликт имен отражён в проблемах; правило один телефон = один человек сохраняется.
                blocking = [i for i in r.issues if not i.startswith("Разные имена")]
                r.status = "Чистая" if not blocking else "Проблема"
                if r.status == "Чистая":
                    clean.append(r)
    return sorted(clean, key=lambda r: r.row)


def load_records(path: Path, sheet_name=None, period=None, infer=True, slash_order="auto"):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        suitable = []
        for sheet in book:
            first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            headers = [tidy(v).casefold() for v in first]
            if all(headers.count(h.casefold()) == 1 for h in HEADERS):
                suitable.append((sheet, [headers.index(h.casefold()) for h in HEADERS]))
        if sheet_name:
            suitable = [item for item in suitable if item[0].title == sheet_name]
        if len(suitable) != 1:
            raise ValueError("Нужен один лист с заголовками: " + ", ".join(HEADERS) +
                             ". Если таких листов несколько, укажите --sheet.")
        sheet, indices = suitable[0]
        rows = [(n, tuple(values[i] if i < len(values) else None for i in indices))
                for n, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2)]
        rows = [(n, raw) for n, raw in rows if any(tidy(v) for v in raw)]
        period_origin = "Указан пользователем" if period else "Без ограничения периода"
        if period is None and infer:
            votes = Counter()
            for _, raw in rows:
                d, issue = resolve_date(raw[2], book.epoch, None, slash_order)
                if d is not None and not issue:
                    votes[(d.year, d.month)] += 1
            if votes:
                common, count = votes.most_common(1)[0]
                if count >= 2 and count * 2 > sum(votes.values()):
                    period = common
                    period_origin = f"Определён по однозначным датам: {count} из {sum(votes.values())}"
        records = []
        for n, raw in rows:
            name = normalize_name(raw[0])
            phone, phone_issue = normalize_phone(raw[1])
            day, date_issue = resolve_date(raw[2], book.epoch, period, slash_order)
            issues = [i for i in ("Нет имени" if not name else "", phone_issue, date_issue) if i]
            records.append(Record(n, raw, name, phone, day, tidy(raw[3]), issues))
        return records, period, period_origin, sheet.title
    finally:
        book.close()


def add_sheet(book, name, headers, rows):
    sheet = book.create_sheet(name)
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    for row in sheet:
        for cell in row:
            # Исходные значения и +7 сохраняются именно текстом, включая строки с '='.
            if isinstance(cell.value, str):
                cell.data_type = "s"
                cell.number_format = "@"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245A81")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        length = max(len(str(c.value or "")) for c in column)
        sheet.column_dimensions[column[0].column_letter].width = min(75, max(16, length + 2))
    return sheet


def write_results(records, clean, period, period_origin, sheet_name, input_path, clean_path, problems_path):
    clean_book = Workbook()
    clean_book.remove(clean_book.active)
    add_sheet(clean_book, "Заявки", HEADERS,
              [(r.name, r.phone, r.day.isoformat(), r.source) for r in clean])
    report = Workbook()
    report.remove(report.active)
    audit_headers = ("Лист", "Строка Excel", "Статус", "Причины", "Выбрана строка Excel",
                     "Исходное имя", "Исходный телефон", "Исходная дата", "Исходный источник",
                     "Имя", "Телефон", "Дата заявки", "Источник")
    def audit_row(r):
        return (sheet_name, r.row, r.status, "; ".join(r.issues), r.kept_row,
                *(v.isoformat() if isinstance(v, (date, datetime)) else text(v) for v in r.raw),
                r.name, r.phone, r.day.isoformat() if r.day else "", r.source)
    problem_rows = [r for r in records if r.issues]
    add_sheet(report, "Проблемы", audit_headers, [audit_row(r) for r in problem_rows])
    add_sheet(report, "Аудит", audit_headers, [audit_row(r) for r in records])
    counts = Counter(r.status for r in records)
    add_sheet(report, "Сводка", ("Показатель", "Значение"), [
        ("Входной файл", input_path.name), ("Лист", sheet_name),
        ("Период", f"{period[0]:04d}-{period[1]:02d}" if period else "Не задан"),
        ("Выбор периода", period_origin), ("Исходных непустых строк", len(records)),
        ("Строк в чистом файле", len(clean)), ("Отклонено дублей", counts["Дубликат"]),
        ("Проблем без включения в чистый файл", counts["Проблема"]),
        ("Строк в списке проблем (включая дубли/конфликты)", len(problem_rows)),
        ("Уникальных распознанных телефонов", len({r.phone for r in records if r.phone})),
    ])
    temp_paths = []
    try:
        for book, destination in ((clean_book, clean_path), (report, problems_path)):
            with tempfile.NamedTemporaryFile(suffix=".xlsx", dir=destination.parent, delete=False) as tmp:
                tmp_path = Path(tmp.name)
            temp_paths.append(tmp_path)
            book.save(tmp_path)
        for temporary, destination in zip(temp_paths, (clean_path, problems_path)):
            temporary.replace(destination)
    finally:
        for path in temp_paths:
            path.unlink(missing_ok=True)
        clean_book.close()
        report.close()
    return counts, len(problem_rows)


def parse_period(s):
    if not re.fullmatch(r"[0-9]{4}-[0-9]{2}", s):
        raise argparse.ArgumentTypeError("Период должен иметь формат YYYY-MM")
    y, m = map(int, s.split("-"))
    try:
        date(y, m, 1)
    except ValueError:
        raise argparse.ArgumentTypeError("Недопустимый год или месяц") from None
    return y, m


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Входной .xlsx")
    parser.add_argument("--out-dir", type=Path, help="Папка результатов (по умолчанию рядом с входным файлом)")
    parser.add_argument("--sheet", help="Имя листа, если листов с заявками несколько")
    choice = parser.add_mutually_exclusive_group()
    choice.add_argument("--period", type=parse_period, help="Ожидаемый месяц YYYY-MM")
    choice.add_argument("--no-infer-period", action="store_true", help="Не определять месяц автоматически")
    parser.add_argument("--slash-order", choices=("auto", "dmy", "mdy"), default="auto",
                        help="Порядок дат через /: auto — по периоду, dmy — день/месяц, mdy — месяц/день")
    parser.add_argument("--force", action="store_true", help="Перезаписать существующие результаты")
    args = parser.parse_args(argv)
    try:
        source = args.input.resolve()
        if source.suffix.lower() != ".xlsx" or not source.is_file():
            raise ValueError("Входной файл не найден или имеет расширение, отличное от .xlsx")
        out = (args.out_dir or source.parent).resolve()
        clean_path, problems_path = (out / f"{source.stem}_{suffix}.xlsx" for suffix in ("clean", "problems"))
        if source in (clean_path, problems_path):
            raise ValueError("Результат не должен совпадать с исходным файлом")
        existing = [p.name for p in (clean_path, problems_path) if p.exists()]
        if existing and not args.force:
            raise ValueError("Результаты уже существуют: " + ", ".join(existing) + ". Укажите --force для перезаписи.")
        records, period, origin, sheet_name = load_records(
            source, args.sheet, args.period, not args.no_infer_period, args.slash_order)
        clean = select_records(records)
        out.mkdir(parents=True, exist_ok=True)
        counts, problem_count = write_results(records, clean, period, origin, sheet_name, source, clean_path, problems_path)
        print(f"Обработано: {len(records)}. Чистых: {len(clean)}. Дублей: {counts['Дубликат']}. "
              f"Проблем без включения в чистый файл: {counts['Проблема']}.")
        print(f"Строк в отчёте проблем: {problem_count}. Полный аудит: {len(records)} строк.")
        print(f"Период: {period[0]:04d}-{period[1]:02d}. {origin}." if period else "Период не определён; неоднозначные даты отмечены в проблемах.")
        print(f"Чистый файл: {clean_path}\nОтчёт проблем: {problems_path}")
        return 0
    except (OSError, ValueError, KeyError, BadZipFile, InvalidFileException) as exc:
        print(f"Ошибка: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
